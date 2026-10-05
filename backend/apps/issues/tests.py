from datetime import timedelta
from unittest import mock

from django.core.management import call_command
from django.test import TestCase, override_settings
from django.utils import timezone
from rest_framework.test import APIClient

from apps.accounts.models import User
from apps.accounts.services import issue_tokens, resolve_login
from apps.cosa.models import Club, ClubMembership
from apps.issues.models import Category, Issue
from apps.notifications.models import Notification


def make_student(email, name="Student"):
    """A student who has already completed the first-login profile."""
    return User.objects.create_user(
        email, full_name=name, roll_number=email.split("@")[0].upper(), branch="CSE",
        batch_year=2023, semester=5, photo="profiles/test.png",
    )


def client_for(user):
    c = APIClient()
    c.credentials(HTTP_AUTHORIZATION=f"Bearer {issue_tokens(user)['access']}")
    return c


class PortalFlowTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_portal", verbosity=0)
        cls.student = make_student("alice@students.iiitr.ac.in", "Alice")
        cls.other_student = make_student("bob@students.iiitr.ac.in", "Bob")
        cls.gensec = User.objects.get(email="gensec_1@students.iiitr.ac.in")
        cls.president = User.objects.get(email="president@iiitr.ac.in")
        cls.mess_sec = User.objects.get(email="messsecretary@iiitr.ac.in")
        cls.sports_sec = User.objects.get(email="sports_secretary@iiitr.ac.in")
        cls.faculty = User.objects.create_user(
            "prof@iiitr.ac.in", full_name="Prof X", user_type=User.UserType.FACULTY
        )
        cls.mess = Category.objects.get(slug="mess")

    def setUp(self):
        self.captured = []

    def raise_issue(self, **extra):
        payload = {"title": "Cold food", "description": "Dinner served cold", "category": self.mess.pk,
                   "priority": "HIGH", **extra}
        with self.captureOnCommitCallbacks(execute=True):
            r = client_for(self.student).post("/api/issues/", payload, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        return r.data

    # ---- auth -------------------------------------------------------------
    def test_login_rules(self):
        self.assertEqual(resolve_login("new@students.iiitr.ac.in").user_type, "STUDENT")
        roll = resolve_login("cs23b1036@iiitr.ac.in")  # roll-number email on the main domain
        self.assertEqual(roll.user_type, "STUDENT")
        self.assertEqual((roll.roll_number, roll.batch_year, roll.branch), ("CS23B1036", 2023, "CSE"))
        self.assertFalse(roll.profile_complete)
        self.assertEqual(resolve_login("hss1@iiitr.ac.in").user_type, "COSA")
        with self.assertRaises(Exception):
            resolve_login("someone@gmail.com")
        with self.assertRaises(Exception):
            resolve_login("unknown.staff@iiitr.ac.in")  # faculty must be pre-added

    def test_refresh_rotation_and_logout(self):
        tokens = issue_tokens(self.student)
        c = APIClient()
        r = c.post("/api/auth/token/refresh/", {"refresh": tokens["refresh"]}, format="json")
        self.assertEqual(r.status_code, 200)
        self.assertIn("refresh", r.data)  # rotated
        # old refresh is blacklisted after rotation
        r2 = c.post("/api/auth/token/refresh/", {"refresh": tokens["refresh"]}, format="json")
        self.assertEqual(r2.status_code, 401)
        # logout blacklists the new one
        c.credentials(HTTP_AUTHORIZATION=f"Bearer {r.data['access']}")
        self.assertEqual(c.post("/api/auth/logout/", {"refresh": r.data["refresh"]}, format="json").status_code, 205)
        self.assertEqual(c.post("/api/auth/token/refresh/", {"refresh": r.data["refresh"]}, format="json").status_code, 401)

    def test_me_and_unauthenticated(self):
        self.assertEqual(APIClient().get("/api/issues/").status_code, 401)
        me = client_for(self.mess_sec).get("/api/auth/me/").data
        self.assertTrue(me["is_cosa"])
        self.assertEqual(me["owned_category"]["name"], "Mess")
        self.assertFalse(me["can_raise_issues"])

    # ---- issues -----------------------------------------------------------
    def test_raise_auto_assigns_and_notifies(self):
        data = self.raise_issue(tagged_members=[self.sports_sec.pk, self.gensec.pk])
        self.assertEqual([a["id"] for a in data["assignees"]], [self.mess_sec.pk])
        self.assertTrue(Notification.objects.filter(recipient=self.mess_sec, kind="ISSUE_CREATED").exists())
        self.assertTrue(Notification.objects.filter(recipient=self.gensec, kind="ISSUE_CREATED").exists())
        self.assertTrue(Notification.objects.filter(recipient=self.sports_sec, kind="ISSUE_TAGGED").exists())
        self.assertFalse(Notification.objects.filter(recipient=self.student).exists())

    def test_first_login_profile(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from io import BytesIO
        from PIL import Image

        new = resolve_login("cs24b1001@iiitr.ac.in")
        c = client_for(new)
        me = c.get("/api/auth/me/").data
        self.assertFalse(me["profile_complete"])
        # can't raise an issue before the profile is complete
        r = c.post("/api/issues/", {"title": "x", "description": "y", "category": self.mess.pk}, format="json")
        self.assertEqual(r.status_code, 403)
        # roll number from the email can't be changed
        r = c.patch("/api/auth/me/", {"roll_number": "XX99"}, format="json")
        self.assertEqual(r.status_code, 400)
        buf = BytesIO()
        Image.new("RGB", (10, 10), "navy").save(buf, "PNG")
        photo = SimpleUploadedFile("me.png", buf.getvalue(), content_type="image/png")
        with self.settings(STORAGES={"default": {"BACKEND": "django.core.files.storage.InMemoryStorage"},
                                     "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"}}):
            r = c.patch("/api/auth/me/", {"full_name": "Krishna", "branch": "CSE", "semester": 3,
                                          "batch_year": 2024, "about": "Hi", "photo": photo}, format="multipart")
            self.assertEqual(r.status_code, 200, r.data)
            self.assertTrue(r.data["profile_complete"])
            self.assertEqual(r.data["year_of_study"], 2)
            self.assertEqual(r.data["branch_label"], "Computer Science and Engineering")
            self.assertTrue(r.data["avatar_url"])
        new.refresh_from_db()
        self.assertIsNotNone(new.profile_completed_at)

    def test_storage_error_is_clean_json(self):
        from botocore.exceptions import ClientError
        from django.core.files.uploadedfile import SimpleUploadedFile
        from io import BytesIO
        from PIL import Image

        buf = BytesIO()
        Image.new("RGB", (10, 10)).save(buf, "PNG")
        photo = SimpleUploadedFile("me.png", buf.getvalue(), content_type="image/png")
        err = ClientError({"Error": {"Code": "AccessDenied", "Message": "denied"}}, "PutObject")
        with mock.patch("django.core.files.storage.FileSystemStorage.save", side_effect=err):
            r = client_for(self.student).patch("/api/auth/me/", {"photo": photo}, format="multipart")
        self.assertEqual(r.status_code, 502)
        self.assertIn("AccessDenied", r.json()["detail"])

    def test_cosa_cannot_raise(self):
        r = client_for(self.gensec).post(
            "/api/issues/", {"title": "x", "description": "y", "category": self.mess.pk}, format="json"
        )
        self.assertEqual(r.status_code, 403)

    def test_no_delete(self):
        issue = self.raise_issue()
        r = client_for(self.gensec).delete(f"/api/issues/{issue['id']}/")
        self.assertEqual(r.status_code, 405)

    def test_assignment_rules(self):
        issue = self.raise_issue()
        url = f"/api/issues/{issue['id']}/assign/"
        # secretary can't assign
        self.assertEqual(client_for(self.mess_sec).post(url, {"assignees": [self.sports_sec.pk]}, format="json").status_code, 403)
        # gen sec and president can
        r = client_for(self.gensec).post(url, {"assignees": [self.mess_sec.pk, self.sports_sec.pk]}, format="json")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(len(r.data["assignees"]), 2)
        r = client_for(self.president).post(
            f"/api/issues/{issue['id']}/faculty/", {"faculty": self.faculty.pk}, format="json"
        )
        self.assertEqual(r.data["faculty"]["id"], self.faculty.pk)
        # faculty can now edit / change status
        r = client_for(self.faculty).post(f"/api/issues/{issue['id']}/status/", {"status": "IN_PROGRESS"}, format="json")
        self.assertEqual(r.status_code, 200)

    def test_status_requires_update_for_blocked_and_completed(self):
        issue = self.raise_issue()
        c = client_for(self.mess_sec)
        url = f"/api/issues/{issue['id']}/status/"
        self.assertEqual(c.post(url, {"status": "BLOCKED"}, format="json").status_code, 400)
        r = c.post(url, {"status": "BLOCKED", "update": "Waiting for vendor"}, format="json")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data["status"], "BLOCKED")
        # other secretary can't
        self.assertEqual(
            client_for(self.sports_sec).post(url, {"status": "IN_PROGRESS"}, format="json").status_code, 403
        )
        # student can't
        self.assertEqual(client_for(self.student).post(url, {"status": "IN_PROGRESS"}, format="json").status_code, 403)

    def test_updates_and_remarks(self):
        issue = self.raise_issue()
        base = f"/api/issues/{issue['id']}"
        self.assertEqual(client_for(self.student).post(f"{base}/updates/", {"body": "hi"}, format="json").status_code, 403)
        self.assertEqual(client_for(self.mess_sec).post(f"{base}/updates/", {"body": "On it"}, format="json").status_code, 201)
        self.assertEqual(client_for(self.mess_sec).post(f"{base}/remarks/", {"body": "x"}, format="json").status_code, 403)
        self.assertEqual(client_for(self.other_student).post(f"{base}/remarks/", {"body": "x"}, format="json").status_code, 403)
        c = client_for(self.student)
        for _ in range(3):
            self.assertEqual(c.post(f"{base}/remarks/", {"body": "Still cold"}, format="json").status_code, 201)
        self.assertEqual(c.post(f"{base}/remarks/", {"body": "4th"}, format="json").status_code, 400)
        self.assertEqual(c.get(f"{base}/remarks/").data["remarks_left_today"], 0)
        timeline = c.get(f"{base}/timeline/").data
        self.assertEqual(len(timeline), 4)

    def test_escalation_after_five_days(self):
        issue = self.raise_issue()
        url = f"/api/issues/{issue['id']}/escalate/"
        c = client_for(self.student)
        self.assertEqual(c.post(url, {"reason": "No progress"}, format="json").status_code, 400)
        Issue.objects.filter(pk=issue["id"]).update(created_at=timezone.now() - timedelta(days=6))
        self.assertEqual(client_for(self.other_student).post(url, {"reason": "x"}, format="json").status_code, 403)
        with self.captureOnCommitCallbacks(execute=True):
            r = c.post(url, {"reason": "No progress"}, format="json")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.data["is_escalated"])
        self.assertTrue(Notification.objects.filter(recipient=self.president, kind="ISSUE_ESCALATED").exists())
        self.assertEqual(c.post(url, {"reason": "again"}, format="json").status_code, 400)  # once per cycle
        esc = client_for(self.president).get("/api/issues/?is_escalated=true").data
        self.assertEqual(esc["count"], 1)

    def test_reopen_restarts_cycle(self):
        issue = self.raise_issue()
        base = f"/api/issues/{issue['id']}"
        client_for(self.mess_sec).post(f"{base}/status/", {"status": "COMPLETED", "update": "Fixed heaters"}, format="json")
        c = client_for(self.student)
        r = c.post(f"{base}/reopen/", {"reason": "Still cold"}, format="json")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["status"], "IN_PROGRESS")
        self.assertEqual(r.data["reopened_count"], 1)
        # can't reopen an open issue
        self.assertEqual(c.post(f"{base}/reopen/", {"reason": "x"}, format="json").status_code, 403)
        # reopen window
        client_for(self.mess_sec).post(f"{base}/status/", {"status": "COMPLETED", "update": "Done"}, format="json")
        Issue.objects.filter(pk=issue["id"]).update(completed_at=timezone.now() - timedelta(days=8))
        self.assertEqual(c.post(f"{base}/reopen/", {"reason": "x"}, format="json").status_code, 403)
        # gen sec can still reopen
        self.assertEqual(client_for(self.gensec).post(f"{base}/reopen/", {"reason": "Audit"}, format="json").status_code, 200)

    def test_upvote_toggle(self):
        issue = self.raise_issue()
        url = f"/api/issues/{issue['id']}/upvote/"
        self.assertEqual(client_for(self.student).post(url).status_code, 403)  # own issue
        c = client_for(self.other_student)
        self.assertEqual(c.post(url).data, {"upvoted": True, "upvote_count": 1})
        self.assertEqual(c.post(url).data, {"upvoted": False, "upvote_count": 0})

    def test_supporters_and_excel_export(self):
        from io import BytesIO
        from openpyxl import load_workbook

        issue = self.raise_issue()
        base = f"/api/issues/{issue['id']}"
        # a student without a complete profile can't add their name
        newbie = User.objects.create_user("cs24b1002@iiitr.ac.in")
        self.assertEqual(client_for(newbie).post(f"{base}/upvote/").status_code, 403)
        client_for(self.other_student).post(f"{base}/upvote/")

        rows = client_for(self.other_student).get(f"{base}/supporters/").data["results"]
        self.assertEqual([(r["full_name"], r["roll_number"], r["role"]) for r in rows],
                         [("Alice", "ALICE", "RAISED"), ("Bob", "BOB", "SUPPORTER")])
        self.assertNotIn("email", rows[0])

        # a private supporter: counted for everyone, named only for COSA and themself
        carol = make_student("carol@students.iiitr.ac.in", "Carol")
        r = client_for(carol).post(f"{base}/support/", {"private": True}, format="json")
        self.assertEqual(r.data, {"my_support": "PRIVATE", "upvote_count": 2})
        seen_by_bob = client_for(self.other_student).get(f"{base}/supporters/").data
        self.assertEqual((seen_by_bob["total"], seen_by_bob["private_hidden"]), (3, 1))
        self.assertNotIn("Carol", [r["full_name"] for r in seen_by_bob["results"]])
        seen_by_cosa = client_for(self.mess_sec).get(f"{base}/supporters/").data
        self.assertEqual(seen_by_cosa["private_hidden"], 0)
        self.assertIn(("Carol", True), [(r["full_name"], r["is_private"]) for r in seen_by_cosa["results"]])
        card = [i for i in client_for(carol).get("/api/issues/").data["results"] if i["id"] == issue["id"]][0]
        self.assertEqual((card["my_support"], card["upvote_count"]), ("PRIVATE", 2))
        # switch to public, then withdraw
        client_for(carol).post(f"{base}/support/", {"private": False}, format="json")
        self.assertIn("Carol", [r["full_name"] for r in client_for(self.other_student).get(f"{base}/supporters/").data["results"]])
        self.assertEqual(client_for(carol).post(f"{base}/unsupport/").data, {"my_support": None, "upvote_count": 1})
        client_for(carol).post(f"{base}/support/", {"private": True}, format="json")

        # only COSA can download
        self.assertEqual(client_for(self.student).get(f"{base}/supporters/export/").status_code, 403)
        r = client_for(self.mess_sec).get(f"{base}/supporters/export/")
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        ws = load_workbook(BytesIO(r.content)).active
        values = [[c for c in row] for row in ws.iter_rows(values_only=True)]
        header = values.index(["S.No", "Name", "Roll No", "Branch", "Year", "Semester", "Email", "Role", "Visibility", "Joined at"])
        self.assertEqual(list(values[header + 1][1:3]), ["Alice", "ALICE"])
        self.assertEqual(list(values[header + 2][1:3]), ["Bob", "BOB"])
        self.assertEqual(values[header + 1][3], "Computer Science and Engineering")
        self.assertEqual((values[header + 3][1], values[header + 3][8]), ("Carol", "Private"))

        # all-issues export respects filters
        r = client_for(self.gensec).get("/api/issues/export/?category=" + str(self.mess.pk))
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(BytesIO(r.content)).active
        titles = [row[1] for row in ws.iter_rows(values_only=True)]
        self.assertIn("Cold food", titles)
        self.assertEqual(client_for(self.student).get("/api/issues/export/").status_code, 403)

    def test_table_view_and_multi_tags(self):
        data = self.raise_issue(tagged_members=[self.sports_sec.pk, self.mess_sec.pk])
        base = f"/api/issues/{data['id']}"
        self.assertEqual(sorted(u["id"] for u in data["tagged_members"]), sorted([self.sports_sec.pk, self.mess_sec.pk]))
        # Only COSA members can be tagged.
        other = make_student("cs23b1099@iiitr.ac.in")
        r = client_for(self.student).post("/api/issues/", {"title": "x", "description": "y", "category": self.mess.pk,
                                                            "tagged_members": [other.pk]}, format="json")
        self.assertEqual(r.status_code, 400)
        client_for(self.mess_sec).post(f"{base}/updates/", {"body": "Talking to vendor"}, format="json")
        client_for(self.student).post(f"{base}/remarks/", {"body": "Still cold"}, format="json")
        r = client_for(self.gensec).get("/api/issues/table/?ordering=-updated_at")
        self.assertEqual(r.status_code, 200)
        row = r.data["results"][0]
        self.assertEqual(row["last_update"]["body"], "Talking to vendor")
        self.assertEqual(row["last_remark"]["body"], "Still cold")
        self.assertEqual(row["created_by"]["roll_number"], self.student.roll_number)
        self.assertEqual(len(row["tagged_members"]), 2)

    def test_board_and_dashboard(self):
        self.raise_issue()
        board = client_for(self.student).get("/api/issues/board/").data
        self.assertEqual([c["status"] for c in board["columns"]], ["NOT_STARTED", "IN_PROGRESS", "BLOCKED", "COMPLETED"])
        self.assertEqual(board["columns"][0]["count"], 1)
        dash = client_for(self.mess_sec).get("/api/cosa/dashboard/").data
        self.assertEqual(dash["my_category"], 1)
        self.assertEqual(dash["assigned_to_me"], 1)
        self.assertEqual(client_for(self.student).get("/api/cosa/dashboard/").status_code, 403)
        mine = client_for(self.student).get("/api/issues/mine/").data
        self.assertFalse(mine["results"][0]["can_escalate"])

    def test_email_privacy(self):
        self.raise_issue()
        card = client_for(self.other_student).get("/api/issues/").data["results"][0]
        self.assertIsNone(card["created_by"]["email"])
        card = client_for(self.gensec).get("/api/issues/").data["results"][0]
        self.assertEqual(card["created_by"]["email"], self.student.email)

    def test_list_query_count_is_constant(self):
        for _ in range(5):
            self.raise_issue()
        c = client_for(self.gensec)
        c.get("/api/issues/")  # warm role cache
        with self.assertNumQueries(5):  # user, count, issues, assignees, tagged
            c.get("/api/issues/")


class CosaTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_portal", verbosity=0)
        cls.student = make_student("alice@students.iiitr.ac.in")
        cls.head = make_student("head@students.iiitr.ac.in")
        cls.cult = User.objects.get(email="cult@students.iiitr.ac.in")
        cls.gensec = User.objects.get(email="gensec_2@students.iiitr.ac.in")

    def test_clubs_events_committees_posts(self):
        cosa = client_for(self.cult)
        r = cosa.post("/api/cosa/clubs/", {"name": "Music Club", "kind": "CULTURAL"}, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        slug = r.data["slug"]
        self.assertEqual(client_for(self.student).post("/api/cosa/clubs/", {"name": "X"}, format="json").status_code, 403)
        cosa.post(f"/api/cosa/clubs/{slug}/members/", {"user_id": self.head.pk, "position": "HEAD"}, format="json")
        club = Club.objects.get(slug=slug)
        self.assertTrue(ClubMembership.objects.filter(club=club, user=self.head, position="HEAD").exists())

        # club head can create an event for their club; a student can't
        start = (timezone.now() + timedelta(days=3)).isoformat()
        with self.captureOnCommitCallbacks(execute=True):
            r = client_for(self.head).post(
                "/api/cosa/events/", {"title": "Open Mic", "club": club.pk, "starts_at": start, "venue": "OAT"}, format="json"
            )
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(r.data["phase"], "UPCOMING")
        self.assertTrue(Notification.objects.filter(recipient=self.student, kind="EVENT").exists())
        self.assertEqual(
            client_for(self.student).post("/api/cosa/events/", {"title": "x", "club": club.pk, "starts_at": start}, format="json").status_code,
            403,
        )
        upcoming = client_for(self.student).get("/api/cosa/events/?when=upcoming").data
        self.assertEqual(upcoming["count"], 1)
        self.assertEqual(client_for(self.student).get(f"/api/cosa/events/?club={slug}").data["count"], 1)

        # committee with open call -> apply -> accept
        until = (timezone.now() + timedelta(days=2)).isoformat()
        r = cosa.post("/api/cosa/committees/", {"name": "Fest Committee", "event": upcoming["results"][0]["id"],
                                                "applications_open_until": until}, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        cid = r.data["id"]
        r = client_for(self.student).post(f"/api/cosa/committees/{cid}/apply/", {"statement": "I can help"}, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(client_for(self.student).post(f"/api/cosa/committees/{cid}/apply/", {"statement": "again"}, format="json").status_code, 400)
        self.assertEqual(client_for(self.student).get(f"/api/cosa/committees/{cid}/applications/").status_code, 403)
        app_id = r.data["id"]
        r = cosa.post(f"/api/cosa/committees/{cid}/applications/{app_id}/decide/", {"decision": "ACCEPTED"}, format="json")
        self.assertEqual(r.status_code, 200, r.data)
        detail = client_for(self.student).get(f"/api/cosa/committees/{cid}/").data
        self.assertEqual(detail["members"][0]["user"]["id"], self.student.pk)
        self.assertEqual(detail["my_application_status"], "ACCEPTED")

        # posts: COSA writes, only gen sec/president pins
        r = cosa.post("/api/cosa/posts/", {"title": "Mess menu revised", "body": "From Monday"}, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        pid = r.data["id"]
        self.assertEqual(cosa.post(f"/api/cosa/posts/{pid}/pin/").status_code, 403)
        self.assertTrue(client_for(self.gensec).post(f"/api/cosa/posts/{pid}/pin/").data["is_pinned"])
        self.assertEqual(client_for(self.gensec).delete(f"/api/cosa/posts/{pid}/").status_code, 405)

    def test_notifications_page(self):
        c = client_for(self.student)
        Notification.objects.create(recipient=self.student, kind="EVENT", title="A")
        Notification.objects.create(recipient=self.student, kind="EVENT", title="B")
        Notification.objects.create(recipient=self.head, kind="EVENT", title="not mine")
        self.assertEqual(c.get("/api/notifications/").data["count"], 2)
        self.assertEqual(c.get("/api/notifications/unread-count/").data["unread"], 2)
        nid = c.get("/api/notifications/").data["results"][0]["id"]
        c.post(f"/api/notifications/{nid}/read/")
        self.assertEqual(c.get("/api/notifications/unread-count/").data["unread"], 1)
        c.post("/api/notifications/read-all/")
        self.assertEqual(c.get("/api/notifications/?is_read=false").data["count"], 0)


@override_settings(GOOGLE_CLIENT_IDS=["test-client.apps.googleusercontent.com"])
class GoogleLoginTests(TestCase):
    def test_google_login_issues_tokens(self):
        claims = {"aud": "test-client.apps.googleusercontent.com", "email": "carol@students.iiitr.ac.in",
                  "email_verified": True, "name": "Carol", "picture": ""}
        with mock.patch("apps.accounts.views.google_id_token.verify_oauth2_token", return_value=claims):
            r = APIClient().post("/api/auth/google/", {"id_token": "x"}, format="json")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertIn("access", r.data)
        self.assertIn("refresh", r.data)
        self.assertEqual(r.data["user"]["user_type"], "STUDENT")

    def test_google_login_rejects_outside_domain(self):
        claims = {"aud": "test-client.apps.googleusercontent.com", "email": "x@gmail.com", "email_verified": True}
        with mock.patch("apps.accounts.views.google_id_token.verify_oauth2_token", return_value=claims):
            r = APIClient().post("/api/auth/google/", {"id_token": "x"}, format="json")
        self.assertEqual(r.status_code, 403)
