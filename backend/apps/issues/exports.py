"""Excel (.xlsx) sheets for COSA: the students on one issue, or all issues."""
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
NAVY = "2B2A7A"


def _local(dt):
    return timezone.localtime(dt).replace(tzinfo=None) if dt else None


def _sheet(ws, title_lines, headers, rows, widths):
    """Title block, a bold navy header row, then data rows with sensible widths."""
    for i, line in enumerate(title_lines, start=1):
        ws.cell(row=i, column=1, value=line).font = Font(bold=(i == 1), size=13 if i == 1 else 11)
    head = len(title_lines) + 2
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=head, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center")
    for r, row in enumerate(rows, start=head + 1):
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            if hasattr(value, "year") and hasattr(value, "hour"):
                cell.number_format = "dd-mmm-yyyy hh:mm"
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = ws.cell(row=head + 1, column=1)
    ws.auto_filter.ref = f"A{head}:{get_column_letter(len(headers))}{head + max(len(rows), 1)}"


def _response(wb, filename):
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type=XLSX)
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


def _student_cols(u):
    return [
        u.full_name,
        u.roll_number,
        u.get_branch_display() if u.branch else "",
        u.year_of_study,
        u.semester,
        u.email,
    ]


def issue_supporters_xlsx(issue):
    """One row per student: the raiser first, then everyone who added their name."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    rows = [[1, *_student_cols(issue.created_by), "First raised", "Public", _local(issue.created_at)]]
    supporters = issue.upvotes.select_related("user").order_by("created_at")
    for n, s in enumerate(supporters, start=2):
        rows.append([n, *_student_cols(s.user), "Also raised", "Private" if s.is_private else "Public",
                     _local(s.created_at)])
    _sheet(
        ws,
        [
            f"Issue #{issue.id}: {issue.title}",
            f"Category: {issue.category.name} · Status: {issue.get_status_display()} · Priority: {issue.get_priority_display()}",
            f"{len(rows)} student(s) · exported {timezone.localtime():%d %b %Y, %I:%M %p}",
        ],
        ["S.No", "Name", "Roll No", "Branch", "Year", "Semester", "Email", "Role", "Visibility", "Joined at"],
        rows,
        [6, 26, 14, 36, 6, 9, 30, 16, 11, 18],
    )
    return _response(wb, f"issue-{issue.id}-students.xlsx")


def all_issues_xlsx(issues):
    """Summary of every issue (filtered list) with how many students are on it."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Issues"
    rows = []
    for i in issues:
        rows.append([
            i.id,
            i.title,
            i.category.name,
            i.get_status_display(),
            i.get_priority_display(),
            "Yes" if i.is_escalated else "",
            i.created_by.full_name,
            i.created_by.roll_number,
            1 + i.upvote_count,
            ", ".join(a.full_name or a.email for a in i.assignees.all()),
            _local(i.created_at),
            _local(i.completed_at),
        ])
    _sheet(
        ws,
        ["COSA Portal · IIIT Raichur — Issues", f"{len(rows)} issue(s) · exported {timezone.localtime():%d %b %Y, %I:%M %p}"],
        ["Issue #", "Title", "Category", "Status", "Priority", "Escalated", "Raised by", "Roll No",
         "Students affected", "Assigned to", "Raised at", "Completed at"],
        rows,
        [9, 40, 22, 14, 10, 10, 24, 14, 10, 30, 18, 18],
    )
    return _response(wb, f"cosa-issues-{timezone.localdate():%Y-%m-%d}.xlsx")
